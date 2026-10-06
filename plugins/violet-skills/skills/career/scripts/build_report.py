"""Turn one job-search run (career/runs/{date}-leads.json) into every output the career skill owes.

Writes, from the memory root:
  career/jobs/{id}.md                         one file per lead (never overwrites an existing file)
  career/tracker.md                           one row per new lead
  career/searches.md                          one dated section per run
  outputs/career-jobs/{date}-job-search.xlsx  workbook: Leads, Scoring, Dropped, Run
  outputs/career-jobs/{date}-job-search.html  report page, published as an artifact

Usage (memory-core .venv has openpyxl):
  .venv/bin/python plugins/violet-skills/skills/career/scripts/build_report.py career/runs/{date}-leads.json
"""
from __future__ import annotations

import html
import json
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

FIT_POINTS = {"Strong": 5, "Possible": 3, "Weak": 1}
BANDS = [(4.2, "Apply now"), (3.5, "Apply"), (3.0, "Consider"), (0.0, "Low priority")]
MARK_WORDS = {"✓": "Meets", "~": "Partly", "✗": "Gap"}


def overall(lead: dict, w: dict) -> float:
    raw = FIT_POINTS[lead["fit"]] * w["fit"] + lead["role_score"] * w["role"] + lead["company_score"] * w["company"]
    return float(Decimal(str(raw)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))  # matches Excel ROUND


def salary_floor(run: dict) -> str:
    """`salary_floor` + `currency`; runs written before the currency field carry `salary_floor_myr`."""
    if "salary_floor_myr" in run:
        return f"MYR {run['salary_floor_myr']:,}"
    return f"{run.get('currency', '')} {run['salary_floor']:,}".strip()


def band(score: float) -> str:
    return next(label for floor, label in BANDS if score >= floor)


def ranked(data: dict) -> list[dict]:
    w = data["run"]["weights"]
    for lead in data["leads"]:
        lead["overall"] = overall(lead, w)
        lead["recommendation"] = band(lead["overall"])
    return sorted(data["leads"], key=lambda l: (l["tier"], -l["overall"]))


# --------------------------------------------------------------------------- markdown

def job_markdown(lead: dict, date: str) -> str:
    rows = "\n".join(f"| {r} | {e} | {m} |" for r, e, m in lead["match"])
    sources = "\n".join(f"- {s}" for s in lead["company_sources"])
    return f"""# {lead['role']} — {lead['company']}

**Status**: found
**Location**: {lead['location']} · {lead['work_mode']} · tier {lead['tier']}
**Source**: {lead['url']}
**Role type**: {lead.get('role_type', 'not set')}
**Seen**: {date} · **Ad captured**: {lead.get('ad_captured', 'summary of the full ad')}
**Salary**: {lead['salary']}
**Posted**: {lead['posted']} · **Applicants**: {lead['applicants']} · **Years asked**: {lead['years_asked']}

## The ad

{lead['ad']}

## Why this job suits you

{lead['why']}

## Fit

| Requirement | Evidence from profile.md | Match |
|-------------|--------------------------|-------|
{rows}

**Gaps**: {lead['gaps']}

**Red flags**: {lead['red_flags']}

**Overall fit**: {lead['fit']}

## Company research

{lead['company_research']}

Sources:
{sources}

## Rating

| Fit | Role | Company | Overall | Recommendation |
|-----|------|---------|---------|----------------|
| {lead['fit']} ({FIT_POINTS[lead['fit']]}) | {lead['role_score']} | {lead['company_score']} | **{lead['overall']} / 5** | {lead['recommendation']} |

**Reason**: {lead['reason']}

## Application

- **CV**:
- **Cover note**:
- **Requirement map**:

## Log

- {date}: found via job search (career skill)
"""


def write_job_files(root: Path, leads: list[dict], date: str) -> list[str]:
    written = []
    for lead in leads:
        path = root / "career" / "jobs" / f"{lead['id']}.md"
        if not path.exists():
            path.write_text(job_markdown(lead, date))
            written.append(path.name)
    return written


def update_tracker(root: Path, leads: list[dict], date: str) -> int:
    path = root / "career" / "tracker.md"
    text = path.read_text()
    added = 0
    for lead in leads:
        if f"jobs/{lead['id']}.md" in text:
            continue
        text = text.rstrip("\n") + (
            f"\n| {lead['role']} | {lead['company']} | {lead['location']} ({lead['tier']}) | "
            f"{lead['fit']} · {lead['overall']} · {lead['recommendation']} | found | {date} | "
            f"Review; tailor CV if applying | [file](jobs/{lead['id']}.md) |"
        )
        added += 1
    path.write_text(text + "\n")
    return added


def update_searches(root: Path, data: dict, leads: list[dict]) -> bool:
    run = data["run"]
    path = root / "career" / "searches.md"
    text = path.read_text()
    snippet_only = sum(1 for l in leads if "snippet" in l.get("ad_captured", ""))
    section = (
        f"## {run['date']}\n\n"
        f"- **Queries**: {'; '.join(run['queries'])}\n"
        f"- **Sources**: {'; '.join(run['sources'])}\n"
        f"- **Role types**: {', '.join(sorted({l.get('role_type', 'not set') for l in leads}))}\n"
        f"- **Leads**: {len(leads) + len(data['dropped'])} found · {len(leads)} shortlisted · "
        f"{len(data['dropped'])} dropped ({'; '.join(f'{d[0]}: {d[2]}' for d in data['dropped'])})\n"
        f"- **Snippet-only ads**: {snippet_only}\n"
        f"- **Notes**: {run['notes']}\n"
        f"- **Outputs**: outputs/career-jobs/{run['date']}-job-search.xlsx · outputs/career-jobs/{run['date']}-job-search.html\n\n"
    )
    marker = "One section per search run, newest first.\n"
    header = f"## {run['date']}\n"
    if header in text:  # re-run of the same date: replace that section
        start = text.index(header)
        nxt = text.find("\n## ", start + len(header))
        text = text[:start] + text[nxt + 1:] if nxt != -1 else text[:start]
    text = text.replace("## YYYY-MM-DD\n\n- **Queries**:\n- **Sources**:\n- **Leads**: N found · N new · N already tracked · N dropped (why)\n", "")
    text = text.replace(marker, marker + "\n" + section, 1) if marker in text else text + "\n" + section
    path.write_text(text.rstrip("\n") + "\n")
    return True  # written or replaced


# --------------------------------------------------------------------------- workbook

def write_workbook(path: Path, data: dict, leads: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    run, w = data["run"], data["run"]["weights"]
    teal, white = "00655F", "FFFFFF"
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head = Font(name="Arial", size=10, bold=True, color=white)
    blue = Font(name="Arial", size=10, color="0000FF")
    link = Font(name="Arial", size=10, color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor=teal)
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D3DBDB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    # Scoring sheet first in code (Leads formulas reference it), shown second.
    sc = wb.active
    sc.title = "Scoring"
    sc_rows = [
        ("Weight: match suitability (fit)", w["fit"], "Share of the overall rating from how well your profile meets the ad."),
        ("Weight: role", w["role"], "Share from the role itself: level, stack, growth, development vs support, pay vs your floor."),
        ("Weight: company", w["company"], "Share from company research: reviews, stability, culture, reputation."),
        ("Weights total (must be 1)", "=SUM(B2:B4)", "Check cell."),
        ("Band: Apply now (overall at least)", 4.2, "Overall rating at or above this is 'Apply now'."),
        ("Band: Apply (overall at least)", 3.5, "At or above this, below 'Apply now'."),
        ("Band: Consider (overall at least)", 3.0, "At or above this, below 'Apply'. Anything lower is 'Low priority'."),
    ]
    sc.append(["Setting", "Value", "Meaning"])
    for r in sc_rows:
        sc.append(list(r))
    sc.append([])
    legend = [
        ("How the overall rating works", ""),
        ("Fit score", "Strong = 5, Possible = 3, Weak = 1, from the Match suitability column on Leads."),
        ("Role and Company scores", "1 to 5, set from the ad and the company research; edit them on Leads (blue cells)."),
        ("Overall rating", "ROUND(Fit score × fit weight + Role score × role weight + Company score × company weight, 1)."),
        ("Cells you can edit", "Blue text on a yellow fill: weights and bands here; Role and Company scores on Leads."),
        ("Source of the scores", f"career skill job search, {run['date']}; reasoning per job is in the 'Reason for rating' column."),
    ]
    for k, v in legend:
        sc.append([k, v])
    for row in sc.iter_rows():
        for c in row:
            c.font, c.alignment = base, wrap
    for c in sc[1]:
        c.font, c.fill = head, head_fill
    for r in (2, 3, 4, 6, 7, 8):
        sc.cell(r, 2).font, sc.cell(r, 2).fill = blue, input_fill
    sc["A10"].font = bold
    sc.column_dimensions["A"].width = 34
    sc.column_dimensions["B"].width = 16
    sc.column_dimensions["C"].width = 90

    # Leads sheet
    ws = wb.create_sheet("Leads", 0)
    headers = [
        "Rank", "Company", "Role", "Location", "Tier", "Work mode", "Posted", "Applicants", "Salary",
        "Years asked", "Why this job suits you", "Match suitability", "Match details", "Gaps", "Red flags",
        "Fit score", "Role score", "Company score", "Overall rating (/5)", "Recommendation",
        "Reason for rating", "Company research", "Company sources", "Job link", "Status",
        "Role type", "Ad captured",
    ]
    ws.append(headers)
    n = len(leads)
    last = n + 1
    for i, lead in enumerate(leads, start=2):
        details = "\n".join(f"{MARK_WORDS[m]}: {r} — {e}" for r, e, m in lead["match"])
        ws.append([
            f"=RANK(S{i},$S$2:$S${last},0)+COUNTIF($S$2:S{i},S{i})-1",
            lead["company"], lead["role"], lead["location"], lead["tier"], lead["work_mode"], lead["posted"],
            lead["applicants"], lead["salary"], lead["years_asked"], lead["why"], lead["fit"], details,
            lead["gaps"], lead["red_flags"],
            f'=IF(L{i}="Strong",5,IF(L{i}="Possible",3,1))',
            lead["role_score"], lead["company_score"],
            f"=ROUND(P{i}*Scoring!$B$2+Q{i}*Scoring!$B$3+R{i}*Scoring!$B$4,1)",
            f'=IF(S{i}>=Scoring!$B$6,"Apply now",IF(S{i}>=Scoring!$B$7,"Apply",IF(S{i}>=Scoring!$B$8,"Consider","Low priority")))',
            lead["reason"], lead["company_research"], "\n".join(lead["company_sources"]), lead["url"], "found",
            lead.get("role_type", ""), lead.get("ad_captured", "summary of the full ad"),
        ])
        ws.cell(i, 24).hyperlink = lead["url"]
    widths = [6, 22, 34, 18, 6, 14, 22, 11, 22, 12, 60, 12, 70, 36, 40, 8, 8, 9, 11, 14, 50, 60, 40, 40, 10, 24, 20]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=1, max_row=last):
        for c in row:
            c.font, c.alignment, c.border = base, wrap, box
    for c in ws[1]:
        c.font, c.fill = head, head_fill
    for i in range(2, last + 1):
        for col in (17, 18):
            ws.cell(i, col).font, ws.cell(i, col).fill = blue, input_fill
        ws.cell(i, 19).font = bold
        ws.cell(i, 24).font = link
    ws.cell(1, 17).comment = Comment("1-5, from the ad: level, stack, growth, development vs support, pay vs floor.", "career skill")
    ws.cell(1, 18).comment = Comment("1-5, from company research: reviews, stability, culture. Sources in column W.", "career skill")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    rng = f"T2:T{last}"
    for text, colour in (("Apply now", "D9EAD3"), ("Apply", "E2F0D9"), ("Consider", "FFF2CC"), ("Low priority", "F4CCCC")):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'T2="{text}"'], fill=PatternFill("solid", fgColor=colour)))

    # Dropped sheet
    dr = wb.create_sheet("Dropped")
    dr.append(["Company", "Role", "Why it was dropped", "Link"])
    for d in data["dropped"]:
        dr.append(d)
    for row in dr.iter_rows():
        for c in row:
            c.font, c.alignment, c.border = base, wrap, box
    for c in dr[1]:
        c.font, c.fill = head, head_fill
    for col, width in zip("ABCD", (26, 40, 60, 60)):
        dr.column_dimensions[col].width = width

    # Run sheet
    rs = wb.create_sheet("Run")
    for k, v in (
        ("Date", run["date"]), ("Profile", run["profile"]),
        ("Location priority", " → ".join(run["location_priority"])),
        ("Salary floor (per month)", salary_floor(run)),
        ("Deal-breakers", "; ".join(run["deal_breakers"])),
        ("Queries", "\n".join(run["queries"])), ("Sources", "\n".join(run["sources"])), ("Notes", run["notes"]),
    ):
        rs.append([k, v])
    for row in rs.iter_rows():
        for c in row:
            c.font, c.alignment = base, wrap
        row[0].font = bold
    rs.column_dimensions["A"].width = 26
    rs.column_dimensions["B"].width = 110

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------- report page

PAGE_CSS = """
:root{--paper:#F2F8F8;--surface:#FFFFFF;--ink:#161C1C;--ink2:#4C5858;--muted:#5A6767;--rule:#D3DBDB;
--accent:#00655F;--soft:#DDF2EF;--warm:#804300;--warm-soft:#FFF3E5;--good:#216439;--good-soft:#E7F3EA;
--bad:#B43E3B;--bad-soft:#FBEDEC;--bar:#20A89B;color-scheme:light}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){color-scheme:dark;--paper:#0C1414;--surface:#121C1C;
--ink:#DCE8E6;--ink2:#B6C0C0;--muted:#94A1A1;--rule:#253434;--accent:#8FCBC5;--soft:#143230;--warm:#E1B48D;--warm-soft:#2E2012;
--good:#9FCAAA;--good-soft:#16281C;--bad:#EEA9A3;--bad-soft:#2E1716}}
:root[data-theme="dark"]{color-scheme:dark;--paper:#0C1414;--surface:#121C1C;--ink:#DCE8E6;--ink2:#B6C0C0;--muted:#94A1A1;
--rule:#253434;--accent:#8FCBC5;--soft:#143230;--warm:#E1B48D;--warm-soft:#2E2012;--good:#9FCAAA;--good-soft:#16281C;--bad:#EEA9A3;--bad-soft:#2E1716}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:15px/1.6 "IBM Plex Sans",-apple-system,"Segoe UI",sans-serif;padding-inline:16px}
.wrap{max-width:980px;margin:0 auto;padding-block:40px 64px}
h1,h2{font-family:"Newsreader",Georgia,serif;font-weight:500;line-height:1.15;text-wrap:balance;margin:0}
h1{font-size:clamp(1.9rem,1.4rem+2vw,2.6rem)}h2{font-size:1.5rem;margin-bottom:8px}
.eyebrow{font:500 12px/1.4 "IBM Plex Mono",ui-monospace,monospace;letter-spacing:.08em;text-transform:uppercase;color:var(--accent);margin:0 0 10px}
.lede{color:var(--ink2);max-width:66ch;margin:10px 0 0}
.summary{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:1px;background:var(--rule);border:1px solid var(--rule);border-radius:2px;margin:24px 0}
.summary div{background:var(--surface);padding:12px 14px}.summary b{display:block;font:500 1.6rem/1 "Newsreader",Georgia,serif;color:var(--accent)}
.summary span{font-size:.8125rem;color:var(--muted)}
.method{background:var(--surface);border:1px solid var(--rule);border-radius:2px;padding:14px 16px;font-size:.9rem;color:var(--ink2);margin-bottom:24px}
.method b{color:var(--ink)}
.filters{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 16px}
.filters button{font:500 .8125rem "IBM Plex Sans",sans-serif;padding:7px 12px;border-radius:999px;border:1px solid var(--rule);background:var(--surface);color:var(--ink2);cursor:pointer}
.filters button[aria-pressed="true"]{background:var(--accent);border-color:var(--accent);color:var(--paper)}
.filters button:focus-visible,a:focus-visible,summary:focus-visible{outline:2px solid var(--accent);outline-offset:2px}
.job{background:var(--surface);border:1px solid var(--rule);border-radius:2px;margin-bottom:12px}
.job>summary{list-style:none;cursor:pointer;display:grid;grid-template-columns:44px minmax(0,1fr) auto;gap:4px 16px;align-items:center;padding:14px 16px}
.job>summary::-webkit-details-marker{display:none}
.rank{font:500 1.4rem/1 "Newsreader",Georgia,serif;color:var(--muted)}
.title b{font-size:1.05rem}.title .co{color:var(--accent);font-weight:600}
.title .meta{display:block;font:400 .8125rem "IBM Plex Mono",monospace;color:var(--muted);margin-top:2px}
.score{text-align:right}.score b{font:500 1.5rem/1 "Newsreader",Georgia,serif}.score span{color:var(--muted);font-size:.8125rem}
.chip{display:inline-block;font:500 .75rem/1 "IBM Plex Mono",monospace;padding:4px 7px;border-radius:2px;margin-top:4px}
.c-apply-now{background:var(--good-soft);color:var(--good)}.c-apply{background:var(--soft);color:var(--accent)}
.c-consider{background:var(--warm-soft);color:var(--warm)}.c-low-priority{background:var(--bad-soft);color:var(--bad)}
.c-type{background:transparent;border:1px solid var(--rule);color:var(--ink2);margin-right:4px}.c-snip{background:var(--warm-soft);color:var(--warm)}
.body{padding:0 16px 16px;border-top:1px solid var(--rule);display:grid;gap:14px}
.body h3{font:600 .8125rem/1.3 "IBM Plex Mono",monospace;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);margin:14px 0 4px}
.body p{margin:0;max-width:72ch}
.bars{display:grid;grid-template-columns:110px minmax(0,1fr) 36px;gap:6px 10px;align-items:center;font-size:.8125rem;max-width:520px}
.bar{height:8px;background:var(--rule);border-radius:2px;overflow:hidden}.bar i{display:block;height:100%;background:var(--bar)}
.bars .v{font-family:"IBM Plex Mono",monospace;text-align:right}
table{border-collapse:collapse;width:100%;font-size:.875rem}th,td{text-align:left;vertical-align:top;padding:7px 8px;border-top:1px solid var(--rule)}
th{font:500 .75rem "IBM Plex Mono",monospace;letter-spacing:.05em;text-transform:uppercase;color:var(--muted);border-top:0}
.tbl{overflow-x:auto}.m{font-weight:600;white-space:nowrap}.m-yes{color:var(--good)}.m-part{color:var(--warm)}.m-no{color:var(--bad)}
.two{display:grid;grid-template-columns:1fr 1fr;gap:16px}.src{font-size:.8125rem;color:var(--muted);overflow-wrap:anywhere}
a{color:var(--accent)}.open{display:inline-block;margin-top:4px;font-weight:600}
.dropped td:nth-child(3){color:var(--ink2)}
footer{margin-top:28px;font-size:.8125rem;color:var(--muted)}
@media (max-width:640px){.job>summary{grid-template-columns:32px minmax(0,1fr);}.score{grid-column:2;text-align:left}.two{grid-template-columns:1fr}}
@media (prefers-reduced-motion:reduce){*{transition:none!important}}
"""

PAGE_JS = """
const buttons=[...document.querySelectorAll('.filters button')];
buttons.forEach(b=>b.addEventListener('click',()=>{
  buttons.forEach(x=>x.setAttribute('aria-pressed',String(x===b)));
  const f=b.dataset.filter;
  document.querySelectorAll('.job').forEach(j=>{j.hidden=!(f==='all'||j.dataset.rec===f);});
}));
"""


def years_label(years: str) -> str:
    return f"{years} yrs" if years[:1].isdigit() else years


def slug(text: str) -> str:
    return text.lower().replace(" ", "-")


def page_html(data: dict, leads: list[dict]) -> str:
    run = data["run"]
    e = html.escape
    counts = {label: sum(1 for l in leads if l["recommendation"] == label) for _, label in BANDS}
    types = sorted({l.get("role_type", "not set") for l in leads})
    cards = []
    for rank, lead in enumerate(leads, start=1):
        match_rows = "".join(
            f'<tr><td>{e(r)}</td><td>{e(ev)}</td><td class="m m-{ {"✓":"yes","~":"part","✗":"no"}[m] }">{MARK_WORDS[m]}</td></tr>'
            for r, ev, m in lead["match"]
        )
        bars = "".join(
            f'<span>{label}</span><span class="bar"><i style="width:{v / 5 * 100:.0f}%"></i></span><span class="v">{v:g}</span>'
            for label, v in (("Match (fit)", FIT_POINTS[lead["fit"]]), ("Role", lead["role_score"]),
                             ("Company", lead["company_score"]), ("Overall", lead["overall"]))
        )
        sources = " · ".join(f'<a href="{e(s)}">{e(s.split("/")[2])}</a>' for s in lead["company_sources"])
        rec = lead["recommendation"]
        cards.append(f"""
<details class="job" data-rec="{slug(rec)}"{" open" if rank <= 3 else ""}>
  <summary>
    <span class="rank">{rank}</span>
    <span class="title"><b>{e(lead['role'])}</b><br><span class="co">{e(lead['company'])}</span>
      <span class="meta">{e(lead['location'])} · {e(lead['work_mode'])} · {e(lead['posted'])} · {e(lead['applicants'])} applicants · {e(years_label(lead['years_asked']))}</span>
      <span class="chip c-type">{e(lead.get('role_type', ''))}</span>{' <span class="chip c-snip">snippet only</span>' if 'snippet' in lead.get('ad_captured', '') else ''}</span>
    <span class="score"><b>{lead['overall']:g}</b> <span>/ 5</span><br><span class="chip c-{slug(rec)}">{e(rec)}</span></span>
  </summary>
  <div class="body">
    <div><h3>Why this job suits you</h3><p>{e(lead['why'])}</p></div>
    <div><h3>Rating</h3><div class="bars">{bars}</div>
      <p style="margin-top:8px"><b>Reason:</b> {e(lead['reason'])}</p></div>
    <div><h3>Match suitability: {e(lead['fit'])}</h3><div class="tbl"><table>
      <tr><th>The ad asks for</th><th>Your evidence</th><th>Match</th></tr>{match_rows}</table></div></div>
    <div class="two">
      <div><h3>Gaps</h3><p>{e(lead['gaps'])}</p></div>
      <div><h3>Red flags</h3><p>{e(lead['red_flags'])}</p></div>
    </div>
    <div><h3>Company research</h3><p>{e(lead['company_research'])}</p><p class="src">Sources: {sources}</p></div>
    <div><h3>Pay</h3><p>{e(lead['salary'])}</p><a class="open" href="{e(lead['url'])}">Open the job ad ↗</a></div>
  </div>
</details>""")
    dropped = "".join(
        f'<tr><td>{e(d[0])}</td><td>{e(d[1])}</td><td>{e(d[2])}</td><td><a href="{e(d[3])}">link</a></td></tr>' for d in data["dropped"]
    )
    filters = '<button type="button" data-filter="all" aria-pressed="true">All ' + str(len(leads)) + "</button>" + "".join(
        f'<button type="button" data-filter="{slug(label)}" aria-pressed="false">{label} {counts[label]}</button>'
        for _, label in BANDS if counts[label]
    )
    w = run["weights"]
    return f"""<title>Job Search {e(run['date'])}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Newsreader:opsz,wght@6..72,500&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap">
<style>{PAGE_CSS}</style>
<div class="wrap">
  <p class="eyebrow">Career · job search · {e(run['date'])}</p>
  <h1>{len(leads)} roles in {e(', '.join(run['location_priority']))}, ranked</h1>
  <p class="lede">Searched against your profile across {len(types)} role types: {e(', '.join(types))}. Floor {e(salary_floor(run))} a month; deal-breakers: {e('; '.join(run['deal_breakers']))}.</p>
  <div class="summary">
    <div><b>{len(leads)}</b><span>leads shortlisted</span></div>
    <div><b>{counts['Apply now']}</b><span>rated Apply now</span></div>
    <div><b>{counts['Apply']}</b><span>rated Apply</span></div>
    <div><b>{len(types)}</b><span>role types</span></div>
    <div><b>{len(data['dropped'])}</b><span>dropped, with reasons</span></div>
  </div>
  <div class="method"><b>How the rating works.</b> Overall = match suitability × {w['fit']:g} + role × {w['role']:g} + company × {w['company']:g}, each out of 5 (fit: Strong 5, Possible 3, Weak 1). Role scores the level, stack, growth and pay against your floor; company scores employee reviews, stability and culture from the research below each job. Apply now ≥ 4.2 · Apply ≥ 3.5 · Consider ≥ 3.0. The same numbers are in the Excel file, where the weights are editable.</div>
  <div class="filters" role="group" aria-label="Filter by recommendation">{filters}</div>
  {''.join(cards)}
  <h2 style="margin-top:32px">Dropped</h2>
  <div class="tbl"><table class="dropped"><tr><th>Company</th><th>Role</th><th>Why</th><th>Ad</th></tr>{dropped}</table></div>
  <footer>{e(run['notes'])} Every fact about you comes from career/profile.md; nothing here was applied for or sent.</footer>
</div>
<script>{PAGE_JS}</script>
"""


def main() -> None:
    leads_file = Path(sys.argv[1]).resolve()
    root = leads_file.parents[2]  # career/runs/x.json -> memory root
    data = json.loads(leads_file.read_text())
    leads = ranked(data)
    date = data["run"]["date"]
    out = root / "outputs" / "career-jobs"
    out.mkdir(parents=True, exist_ok=True)
    written = write_job_files(root, leads, date)
    added = update_tracker(root, leads, date)
    logged = update_searches(root, data, leads)
    write_workbook(out / f"{date}-job-search.xlsx", data, leads)
    (out / f"{date}-job-search.html").write_text(page_html(data, leads))
    print(f"job files: {len(written)} · tracker rows: {added} · search log: {'written' if logged else 'unchanged'}")
    for i, l in enumerate(leads, start=1):
        print(f"{i:>2}. {l['overall']:.1f} {l['recommendation']:<12} {l['fit']:<8} {l['company']} — {l['role']}")
    print(f"wrote {out / (date + '-job-search.xlsx')}")
    print(f"wrote {out / (date + '-job-search.html')}")


if __name__ == "__main__":
    main()
